"""Tự kiểm tra không cần API: tạo pptx/docx/xlsx mẫu, thử cơ chế bắt lỗi sandbox, dựng giao diện offscreen.
Chạy:  .venv\\Scripts\\python selftest.py"""
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / "src"))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import tools  # noqa: E402
from pptx_design_engine import fit_font_size  # noqa: E402

# chống tràn: chữ dài phải nhỏ hơn chữ ngắn trong cùng khung
assert fit_font_size("Ngắn", 4, 1, 24) == 24
assert fit_font_size("Rất dài " * 80, 4, 1, 24) < 14

from PIL import Image  # noqa: E402
img = tools.IMAGES_DIR / "selftest.jpg"
Image.new("RGB", (900, 600), (120, 170, 140)).save(img)

code = f"""
deck = Deck(theme="medical_teal", footer="Trà xanh")
deck.hero("Trà xanh và sức khoẻ người cao tuổi", "Lợi ích, cách dùng và lưu ý", image=r"{img}")
deck.split("Trà xanh có gì?", [("Chất chống oxy hoá", "Catechin giúp bảo vệ tế bào."), ("Caffeine nhẹ", "Tỉnh táo, ít gây hồi hộp hơn cà phê.")], r"{img}")
deck.cards("Lợi ích chính", [("Tim mạch", "Hỗ trợ giảm cholesterol."), ("Trí nhớ", "Giúp tỉnh táo."), ("Tiêu hoá", "Dễ chịu sau bữa ăn."), ("Răng miệng", "Giảm vi khuẩn.")])
deck.chart("Lượng catechin", "bar", ["Trà xanh", "Trà ô long", "Trà đen"], [120, 60, 25], "Trà xanh nhiều nhất.", unit=" mg")
deck.timeline("Cách pha", [("Đun nước", "80 độ"), ("Tráng ấm", "Nước nóng"), ("Ngâm", "2-3 phút")])
deck.statement("Uống vừa phải, 2–3 chén mỗi ngày.", "Hỏi bác sĩ nếu đang dùng thuốc")
deck.closing()
deck.save("selftest_slides.pptx")
doc = new_document("emerald_cream"); add_title(doc, "Báo cáo", "Thử nghiệm")
doc.add_heading("Mục 1", level=1); doc.add_paragraph("Nội dung tiếng Việt có dấu.")
add_table(doc, ["Món", "Giá"], [["Trà", 50000], ["Bánh", 20000]]); add_chart_to_doc(doc, "pie", ["A", "B"], [3, 2])
save_document(doc, "selftest_doc.docx")
wb = new_workbook(); ws = wb.active
write_table(ws, "CHI TIÊU THÁNG", ["Ngày", "Khoản", "Số tiền"], [["01/10", "Chợ", 350000], ["02/10", "Điện", 820000]], money_cols=["Số tiền"])
save_workbook(wb, "selftest_sheet.xlsx")
"""
ok, out, err, files = tools.execute_python(code)
assert ok, err
assert {Path(f).suffix for f in files} == {".pptx", ".docx", ".xlsx"}, files
print("Tạo file OK:", *files, sep="\n  ")

from openpyxl import load_workbook  # noqa: E402
ws = load_workbook(next(f for f in files if f.endswith(".xlsx"))).active
assert ws["C6"].value == "=SUM(C4:C5)", ws["C6"].value

ok, _, err, _ = tools.execute_python("raise ValueError('lỗi thử')")
assert not ok and "ValueError" in err
ok, _, _, _ = tools.execute_python("import sys; sys.exit(1)")
assert not ok  # code AI không được tắt app
print("Bắt lỗi sandbox OK")

from agent import MAX_FIX, Callbacks, _run_tool  # noqa: E402
state = {"fails": 0}
for _ in range(MAX_FIX + 1):
    r = _run_tool("execute_python_file_processor", {"code": "1/0"}, {}, Callbacks(), state)
    assert "ZeroDivisionError" in r["traceback"]
assert "hết lượt" in r["note"]
r = _run_tool("execute_python_file_processor", {"code": "1/0"}, {}, Callbacks(), state)
assert "traceback" not in r  # quá 3 lần sửa thì không chạy nữa
print("Giới hạn tự sửa lỗi OK")

# Đọc file văn phòng: header/footer + bảng đúng vị trí, Excel dài giữ dòng tổng + công thức, CSV cp1258, TXT UTF-16
import docx as _docx  # noqa: E402
import openpyxl as _xl  # noqa: E402
d = _docx.Document()
d.sections[0].header.paragraphs[0].text = "ĐẦU TRANG XYZ"
d.add_paragraph("Trước bảng")
d.add_table(rows=1, cols=2).rows[0].cells[0].text = "Ô BẢNG"
d.add_paragraph("Sau bảng")
d.save(tools.OUTPUT_DIR / "_t.docx")
txt = tools.extract_file_text(str(tools.OUTPUT_DIR / "_t.docx"))
assert "ĐẦU TRANG XYZ" in txt and txt.index("Trước bảng") < txt.index("Ô BẢNG") < txt.index("Sau bảng"), txt
wb = _xl.Workbook()
for i in range(1, 501):
    wb.active.append([f"Dòng {i}", i])
wb.active.append(["Tổng cộng", "=SUM(B1:B500)"])
wb.save(tools.OUTPUT_DIR / "_t.xlsx")
txt = tools.extract_file_text(str(tools.OUTPUT_DIR / "_t.xlsx"))
assert "Tổng cộng" in txt and "=SUM(B1:B500)" in txt and "501 dòng" in txt, txt[-300:]
import unicodedata as _ud  # noqa: E402
_out = []  # cp1258: chữ â ê ô ă ơ ư liền, chỉ dấu thanh tách rời
for _ch in _ud.normalize("NFD", "Họ tên,Nội"):
    if _ch in "̛̂̆":
        _j = max(i for i, c in enumerate(_out) if not _ud.combining(c))
        _out[_j] = _ud.normalize("NFC", _out[_j] + _ch)
    else:
        _out.append(_ch)
(tools.OUTPUT_DIR / "_t.csv").write_bytes("".join(_out).encode("cp1258"))
assert "Họ tên" in tools.extract_file_text(str(tools.OUTPUT_DIR / "_t.csv"))
(tools.OUTPUT_DIR / "_t.txt").write_text("Ghi chú tiếng Việt", encoding="utf-16")
assert tools.extract_file_text(str(tools.OUTPUT_DIR / "_t.txt")) == "Ghi chú tiếng Việt"
for ext in ("docx", "xlsx", "csv", "txt"):
    (tools.OUTPUT_DIR / f"_t.{ext}").unlink()
print("Đọc file văn phòng OK")
ok, _, err, _ = tools.execute_python("Document().add_paragraph().add_run('x').font.color.rgb = RGBColor(1, 2, 3)\n"
                                     "p = Presentation(); s = p.slides.add_slide(p.slide_layouts[6])\n"
                                     "b = s.shapes.add_textbox(0, 0, 9, 9); b.text_frame.text = 'x'\n"
                                     "b.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(1, 2, 3)")
assert ok, err  # RGBColor trong sandbox phải dùng được cho cả Word và PowerPoint

import render  # noqa: E402  hiển thị: công thức thành ảnh, tiền không bị nhầm là công thức, code giữ nguyên để sao chép
_h, _codes = render.to_html("Ta có \\(a^2+b^2\\), $$\\frac{1}{2}$$, giá 5$ và 10$.\n\n```python\nx = 1\n```\n\n| A | B |\n|---|---|\n| 1 | 2 |",
                            render.palette("dark"))
assert _h.count("<img") == 2 and "5$ và 10$" in _h and _codes == ["x = 1\n"] and "<th" in _h, _h[:300]
print("Hiển thị Markdown / LaTeX / code OK")

# Đợt sửa theo báo cáo QA 29/10: im lặng không gửi API, giữ số lẻ, không cộng cột đơn giá, tổng cột A, ngân sách web
import io as _io  # noqa: E402
import wave as _wave  # noqa: E402
import math as _math  # noqa: E402
import struct as _struct  # noqa: E402


def _wav(samples):
    b = _io.BytesIO()
    with _wave.open(b, "wb") as w_:
        w_.setnchannels(1); w_.setsampwidth(2); w_.setframerate(16000)
        w_.writeframes(b"".join(_struct.pack("<h", s) for s in samples))
    return b.getvalue()


assert not tools.has_speech(_wav([0] * 32000)) and tools.transcribe(_wav([0] * 32000), {}) == ""
assert tools.has_speech(_wav([int(6000 * _math.sin(i / 8)) for i in range(16000)]))
import office_helpers as _oh  # noqa: E402
assert _oh.vn_number(1234.56) == "1.234,56" and _oh.vn_number(1200000) == "1.200.000"
_wb = _oh.new_workbook()
_pos = _oh.write_table(_wb.active, "BG", ["Hàng", "Đơn giá", "Thành tiền"], [["A", 100, 200], ["B", 50, 150]],
                       money_cols=["Đơn giá", "Thành tiền"])
assert _wb.active.cell(_pos["total_row"], 2).value is None and _wb.active.cell(_pos["total_row"], 3).value.startswith("=SUM")
_wb = _oh.new_workbook()
_pos = _oh.write_table(_wb.active, "T", ["Tiền", "Ghi chú"], [[10, "a"], [20, "b"]], money_cols=["Tiền"])
assert _wb.active.cell(_pos["total_row"], 1).value.startswith("=SUM") and _wb.active.cell(_pos["total_row"], 2).value == "Tổng cộng"
from agent import MAX_WEB_CALLS  # noqa: E402
_st = {"fails": 0, "web": MAX_WEB_CALLS}
assert "hết" in _run_tool("web_search", {"query": "x"}, {}, Callbacks(), _st)["error"]
import voice as _voice  # noqa: E402
assert "Lỗi:" in _voice.transcript_markdown({"title": "t", "display": [{"role": "assistant", "text": "", "error": "Hết hạn"}]})
print("Các bản sửa theo báo cáo QA OK")

import shortcuts  # noqa: E402  lối tắt: không ghi đè lối tắt "Agenty" của app khác
_T = tools.OUTPUT_DIR / "_lnk"
_T.mkdir(exist_ok=True)
_sh = shortcuts._shell()
shortcuts.write_link(_T / "Agenty.lnk", Path(r"C:\Windows\notepad.exe"))
assert shortcuts.choose_name(_sh, [_T], Path(sys.executable)) == "Agenty (2)"
_mine = shortcuts._place(_sh, _T, Path(sys.executable), "Agenty (2)")
assert shortcuts.choose_name(_sh, [_T], Path(sys.executable)) == "Agenty (2)"  # lần sau giữ tên cũ
assert shortcuts._target(_sh, _T / "Agenty.lnk").lower().endswith("notepad.exe")  # không đụng lối tắt app khác
import shutil as _sh_util  # noqa: E402
_sh_util.rmtree(_T)
print("Lối tắt OK")

assert tools.resolve_path("Desktop").name in ("Desktop",) and tools.resolve_path("Desktop/a.docx").name == "a.docx"
listing = tools.list_directory(str(tools.OUTPUT_DIR), "*.pptx")
assert any(i["name"] == "selftest_slides.pptx" for i in listing["items"]), listing
assert "error" in tools.list_directory("Z:/khong_co_thu_muc_nay")
print("Truy cập file trong máy OK")

for f in files:
    Path(f).unlink()

from PySide6.QtWidgets import QApplication  # noqa: E402
from ui import MainWindow  # noqa: E402
app = QApplication([])
from ui import apply_theme; apply_theme(app); w = MainWindow()
w.show()
app.processEvents()
w.grab().save(str(Path(__file__).parent / "screenshot.png"))
print("Dựng giao diện OK -> screenshot.png")
