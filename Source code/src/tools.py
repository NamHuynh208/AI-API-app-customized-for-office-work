"""Công cụ cho AI: tìm web, đọc trang, tải ảnh, tạo ảnh AI, chạy Python tạo file, đọc file đính kèm."""
import base64
import contextlib
import io
import re
import time
import traceback
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import httpx
from bs4 import BeautifulSoup
from PIL import Image, ImageDraw, ImageOps

from config import IMAGES_DIR, OUTPUT_DIR
from llm import make_client

UA = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/130.0 Safari/537.36"
}


def _ddgs_search(kind: str, query: str, max_results: int) -> list:
    try:
        from ddgs import DDGS
    except ImportError:
        from duckduckgo_search import DDGS
    # Các nguồn hay chặn tạm thời khi gọi dồn dập: lần sau đổi sang nhóm nguồn khác
    backends = ["auto", "google, bing, yahoo"] if kind == "text" else ["auto", "auto"]
    for i, backend in enumerate(backends):
        try:
            rows = getattr(DDGS(), kind)(query, max_results=max_results, backend=backend)
            if rows:
                return rows
        except Exception:
            pass
        if i == 0:
            time.sleep(1.5)
    return []


# ---------- Tìm kiếm ----------

def web_search(query: str, max_results: int = 6) -> dict:
    rows = _ddgs_search("text", query, int(max_results))
    if not rows:
        return {"results": [], "note": "Không tìm thấy kết quả. Thử từ khoá khác (ngắn hơn hoặc tiếng Anh)."}
    return {"results": [{"title": r.get("title"), "url": r.get("href"), "snippet": r.get("body")} for r in rows]}


def read_webpage(url: str, max_chars: int = 12000) -> dict:
    r = httpx.get(url, headers=UA, timeout=20, follow_redirects=True)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "svg"]):
        tag.decompose()
    lines = (line.strip() for line in soup.get_text("\n").splitlines())
    text = "\n".join(line for line in lines if line)
    title = soup.title.get_text(strip=True) if soup.title else ""
    return {"url": url, "title": title, "content": text[:max_chars]}


# ---------- Ảnh ----------

def save_image_bytes(data: bytes, prefix: str, min_side: int = 0) -> str | None:
    """Kiểm tra ảnh hợp lệ bằng Pillow rồi chuẩn hoá sang .jpg (hoặc .png nếu có nền trong suốt)."""
    Image.open(io.BytesIO(data)).verify()
    img = ImageOps.exif_transpose(Image.open(io.BytesIO(data)))
    if min(img.size) < min_side:
        return None
    alpha = img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info)
    name = f"{prefix}_{uuid.uuid4().hex[:8]}"
    if alpha:
        path = IMAGES_DIR / f"{name}.png"
        img.convert("RGBA").save(path)
    else:
        path = IMAGES_DIR / f"{name}.jpg"
        img.convert("RGB").save(path, quality=90)
    return str(path)


def search_and_download_images(query: str, max_images: int = 4) -> dict:
    max_images = max(1, min(int(max_images), 8))
    results = _ddgs_search("images", query, max_images * 4)
    urls = [r["image"] for r in results if r.get("image")]

    def fetch(url):
        try:
            r = httpx.get(url, headers=UA, timeout=12, follow_redirects=True)
            r.raise_for_status()
            return save_image_bytes(r.content, "anh", min_side=300)
        except Exception:
            return None

    with ThreadPoolExecutor(6) as ex:
        paths = [p for p in ex.map(fetch, urls) if p]
    for extra in paths[max_images:]:
        Path(extra).unlink(missing_ok=True)
    return {"local_image_paths": paths[:max_images]}


def _image_from_chat(client, model, prompt, image_path=None) -> bytes | None:
    """Fallback: model ảnh chỉ chạy qua chat completions (trả data URL / link ảnh / trường images).
    Có image_path thì gửi kèm ảnh gốc để model sửa."""
    content = prompt if not image_path else [
        {"type": "text", "text": prompt}, {"type": "image_url", "image_url": {"url": image_to_data_url(image_path, 1600)}}]
    msg = client.chat.completions.create(model=model, messages=[{"role": "user", "content": content}]).choices[0].message
    candidates = [msg.content or ""]
    for item in (msg.model_extra or {}).get("images") or []:
        candidates.append(str((item.get("image_url") or {}).get("url", "")) if isinstance(item, dict) else str(item))
    for text in candidates:
        m = re.search(r"data:image/[\w+]+;base64,([A-Za-z0-9+/=]+)", text)
        if m:
            return base64.b64decode(m.group(1))
        m = re.search(r"https?://\S+?\.(?:png|jpe?g|webp)(?:\?\S*)?(?=[\s)\"']|$)", text) or re.search(r"!\[[^\]]*\]\((https?://[^)]+)\)", text)
        if m:
            url = m.group(1) if m.groups() else m.group(0)
            return httpx.get(url, timeout=60, follow_redirects=True).content
    return None


def generate_ai_image(prompt: str, model: str | None = None, cfg: dict | None = None) -> dict:
    client = make_client(cfg)
    models = list(dict.fromkeys(m for m in (model, cfg.get("default_image_model"), "gpt-image-2.5", "gemini-3.1-flash-image") if m))
    errors = []
    for m in models:
        try:
            item = client.images.generate(model=m, prompt=prompt, n=1, size="1536x1024").data[0]
            data = base64.b64decode(item.b64_json) if item.b64_json else httpx.get(item.url, timeout=60, follow_redirects=True).content
            return {"local_image_path": save_image_bytes(data, "ai"), "model": m}
        except Exception as e:
            errors.append(f"{m} (images): {e}")
    for m in models:
        try:
            data = _image_from_chat(client, m, prompt)
            if data:
                return {"local_image_path": save_image_bytes(data, "ai"), "model": m}
        except Exception as e:
            errors.append(f"{m} (chat): {e}")
    return {"error": "Không tạo được ảnh. " + " | ".join(errors)[:1500]}


def edit_image(image_path: str, prompt: str, cfg: dict) -> dict:
    """Sửa ảnh có sẵn bằng AI (thêm / bỏ chi tiết, đổi nền, đổi phong cách...). Ảnh gốc giữ nguyên, lưu bản mới."""
    client = make_client(cfg)
    src = Path(image_path)
    png = io.BytesIO()
    ImageOps.exif_transpose(Image.open(src)).convert("RGBA").save(png, "PNG")
    errors = []
    for m in dict.fromkeys(x for x in (cfg.get("default_image_model"), "gpt-image-2.5") if x and "gemini" not in x):
        try:
            item = client.images.edit(model=m, image=(src.stem + ".png", png.getvalue(), "image/png"), prompt=prompt).data[0]
            data = base64.b64decode(item.b64_json) if item.b64_json else httpx.get(item.url, timeout=60, follow_redirects=True).content
            return {"local_image_path": save_image_bytes(data, "anh_sua"), "model": m}
        except Exception as e:
            errors.append(f"{m}: {e}")
    try:  # dự phòng: model gemini sửa ảnh qua chat
        data = _image_from_chat(client, "gemini-3.1-flash-image", f"Edit this image: {prompt}. Return the edited image.", str(src))
        if data:
            return {"local_image_path": save_image_bytes(data, "anh_sua"), "model": "gemini-3.1-flash-image"}
    except Exception as e:
        errors.append(f"gemini: {e}")
    return {"error": "Không sửa được ảnh. " + " | ".join(errors)[:1500]}


SILENCE_RMS = 150  # ngưỡng năng lượng PCM16; dưới mức này coi như không có tiếng nói


def has_speech(wav: bytes) -> bool:
    """Kiểm tra năng lượng âm thanh theo từng khung 30 ms: cần ít nhất ~0,3 giây có tiếng."""
    import wave
    import numpy as np
    with wave.open(io.BytesIO(wav)) as w:
        if w.getsampwidth() != 2:
            return True  # định dạng lạ: để model tự xử lý
        rate = w.getframerate()
        a = np.frombuffer(w.readframes(w.getnframes()), dtype=np.int16).astype(np.float32)
        if w.getnchannels() > 1:
            a = a[: len(a) // w.getnchannels() * w.getnchannels()].reshape(-1, w.getnchannels()).mean(axis=1)
    frame = max(int(rate * 0.03), 1)
    n = len(a) // frame
    if n == 0:
        return False
    rms = np.sqrt((a[: n * frame].reshape(n, frame) ** 2).mean(axis=1))
    return int((rms > SILENCE_RMS).sum()) * 0.03 >= 0.3


def transcribe(wav: bytes, cfg: dict) -> str:
    """Giọng nói → chữ. API không có endpoint chép lời nên nhờ model nghe được âm thanh (gemini) chép nguyên văn.
    Đoạn im lặng trả "" mà không gọi API."""
    if not has_speech(wav):
        return ""
    client = make_client(cfg)
    r = client.chat.completions.create(model=cfg.get("voice_model", "gemini-3.8-flash-high"), messages=[{"role": "user", "content": [
        {"type": "text", "text": "Chép lại NGUYÊN VĂN lời nói trong đoạn ghi âm (thường là tiếng Việt, có dấu đầy đủ). "
                                 "Chỉ trả về đúng lời nói, không giải thích, không thêm dấu ngoặc. Không có lời nói thì trả về chuỗi rỗng."},
        {"type": "input_audio", "input_audio": {"data": base64.b64encode(wav).decode(), "format": "wav"}}]}])
    return (r.choices[0].message.content or "").strip().strip('"')


# ---------- Python sandbox ----------

_PRELUDE = """
import os, json, math, datetime
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import docx
from docx import Document
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
import pandas as pd
import xlrd
import reportlab
from pypdf import PdfReader, PdfWriter
import docx.shared as _ds
import pptx.dml.color as _pc
class RGBColor(_pc.RGBColor, _ds.RGBColor):  # một màu dùng được cho cả Word lẫn PowerPoint
    def __new__(cls, r, g, b):
        return tuple.__new__(cls, (r, g, b))
from docx.shared import Cm
from pptx_design_engine import *
from office_helpers import *
"""


def _snapshot() -> dict:
    out = {}
    for p in OUTPUT_DIR.rglob("*"):
        if p.is_file() and IMAGES_DIR not in p.parents:
            with contextlib.suppress(OSError):
                out[p] = p.stat().st_mtime
    return out


def execute_python(code: str) -> tuple[bool, str, str, list[str]]:
    """Chạy code AI viết. Trả (thành_công, stdout, traceback, các file mới/được sửa trong OUTPUT_DIR)."""
    ns = {"__name__": "__sandbox__", "OUTPUT_DIR": OUTPUT_DIR, "IMAGES_DIR": IMAGES_DIR}
    before = _snapshot()
    out = io.StringIO()
    try:
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(out):
            exec(_PRELUDE, ns)
            exec(compile(code, "<ai_code>", "exec"), ns)
        ok, err = True, ""
    except BaseException:  # bắt cả SystemExit để code AI không tắt được app
        ok, err = False, traceback.format_exc(limit=8)
    finally:
        with contextlib.suppress(Exception):
            import matplotlib.pyplot as plt
            plt.close("all")
    files = [str(p) for p, t in _snapshot().items() if before.get(p) != t]
    return ok, out.getvalue()[-4000:], err[-5000:], files


# ---------- Xem trước slide ----------

def render_presentation(path: str, max_slides: int = 16) -> list[str]:
    """Xuất từng slide ra PNG bằng PowerPoint hoặc WPS (qua COM) rồi ghép thành ảnh lưới 2 cột.
    Trả danh sách đường dẫn ảnh lưới (mỗi ảnh tối đa 6 slide)."""
    import os
    import pythoncom
    import win32com.client

    src = os.path.abspath(path)
    out = IMAGES_DIR / "preview" / uuid.uuid4().hex[:8]
    out.mkdir(parents=True, exist_ok=True)
    pythoncom.CoInitialize()  # code AI chạy ở luồng nền
    try:
        app = win32com.client.dynamic.Dispatch("PowerPoint.Application")  # WPS cũng đăng ký tên này
        pres = app.Presentations.Open(src)  # WPS chỉ nhận một tham số
        n = min(int(pres.Slides.Count), max_slides)
        for i in range(1, n + 1):
            pres.Slides(i).Export(str(out / f"s{i:02d}.png"), "PNG", 1280, 720)
        pres.Close()
        with contextlib.suppress(Exception):
            if int(app.Presentations.Count) == 0:  # không tắt nếu người dùng đang mở file khác
                app.Quit()
    finally:
        pythoncom.CoUninitialize()
    files = sorted(out.glob("s*.png"))
    sheets = []
    for start in range(0, len(files), 6):
        group = files[start:start + 6]
        sheet = Image.new("RGB", (2 * 660 + 10, ((len(group) + 1) // 2) * 400 + 10), "#888888")
        for k, f in enumerate(group):
            th = Image.open(f).convert("RGB").resize((640, 360))
            x, y = 10 + (k % 2) * 660, 10 + (k // 2) * 400
            sheet.paste(th, (x, y + 30))
            ImageDraw.Draw(sheet).text((x, y + 8), f"Slide {start + k + 1}", fill="white")
        p = out / f"sheet{start // 6 + 1}.jpg"
        sheet.save(p, quality=85)
        sheets.append(str(p))
    return sheets


def preview_images(path: str, pages: list[int] | None = None) -> list[str]:
    """Ảnh xem trước: ảnh → chính ảnh đó; slide → ảnh lưới; pdf → các trang chọn (mặc định 4 trang đầu);
    Word/Excel → chuyển pdf rồi dựng ảnh."""
    ext = Path(path).suffix.lower()
    if ext in IMAGE_EXTS:
        Image.open(path).verify()
        return [path]
    if ext in (".pptx", ".ppt"):
        return render_presentation(path)
    if ext != ".pdf":
        from office_helpers import convert_file
        path = convert_file(path, "pdf", out_dir=str(OUTPUT_DIR.parent / "converted"))
    return pdf_page_images(path, pages=pages or [1, 2, 3, 4], scale=1.3)


def parse_pages(spec) -> list[int] | None:
    """'7' | '2-4,9' | [7, 8] → [7] | [2, 3, 4, 9] | [7, 8]; rỗng → None."""
    if not spec:
        return None
    if isinstance(spec, (list, tuple)):
        return [int(x) for x in spec]
    out = []
    for part in str(spec).replace(" ", "").split(","):
        if "-" in part:
            a, b = part.split("-", 1)
            out += list(range(int(a), int(b) + 1))
        elif part:
            out.append(int(part))
    return out[:12]


# ---------- File trong máy ----------

KNOWN_FOLDERS = {  # FOLDERID của Windows
    "Desktop": "{B4BFCC3A-DB2C-424C-B029-7FE99A87C641}", "Documents": "{FDD39AD0-238F-46AF-ADB4-6C85480369C7}",
    "Downloads": "{374DE290-123F-4565-9164-39C4925E467B}", "Pictures": "{33E28130-4E1E-4676-835A-98395C3BC3BB}",
    "Music": "{4BD8D571-6D19-48D3-BE97-422220080E43}", "Videos": "{18989B1D-99B5-455B-841C-AB7C74E4DDFC}"}


def _known_folder(guid: str) -> Path | None:
    """Đường dẫn thật của thư mục Desktop/Documents... (kể cả khi OneDrive chuyển chỗ)."""
    import ctypes
    import uuid as _uuid
    from ctypes import wintypes
    try:
        class GUID(ctypes.Structure):
            _fields_ = [("d1", wintypes.DWORD), ("d2", wintypes.WORD), ("d3", wintypes.WORD), ("d4", ctypes.c_ubyte * 8)]
        u = _uuid.UUID(guid)
        g = GUID(u.time_low, u.time_mid, u.time_hi_version, (ctypes.c_ubyte * 8)(*u.bytes[8:]))
        out = ctypes.c_wchar_p()
        if ctypes.windll.shell32.SHGetKnownFolderPath(ctypes.byref(g), 0, None, ctypes.byref(out)) != 0:
            return None
        path = Path(out.value)
        ctypes.windll.ole32.CoTaskMemFree(out)
        return path
    except Exception:
        return None


def _known_dirs() -> dict:
    home = Path.home()
    out = {"Home": home}
    for name, guid in KNOWN_FOLDERS.items():
        out[name] = _known_folder(guid) or home / name
    return out


def user_dirs() -> str:
    return ", ".join(f"{k} = {v}" for k, v in _known_dirs().items())


def resolve_path(path: str) -> Path:
    """Nhận 'Desktop', 'Desktop/abc.docx', '~/...' hoặc đường dẫn đầy đủ."""
    path = str(path).strip().strip('"')
    head, _, rest = path.replace("\\", "/").partition("/")
    base = _known_dirs().get(head.capitalize()) if head and ":" not in head else None
    p = base / rest if base else Path(path).expanduser()
    return p.resolve()


def list_directory(path: str, pattern: str = "*", limit: int = 200) -> dict:
    p = resolve_path(path)
    if not p.is_dir():
        return {"error": f"Không thấy thư mục {p}"}
    items = []
    for child in sorted(p.glob(pattern or "*"), key=lambda c: (not c.is_dir(), c.name.lower()))[:limit]:
        with contextlib.suppress(OSError):
            st = child.stat()
            items.append({"name": child.name, "type": "folder" if child.is_dir() else "file",
                          "size": None if child.is_dir() else st.st_size,
                          "modified": time.strftime("%d/%m/%Y %H:%M", time.localtime(st.st_mtime))})
    return {"path": str(p), "items": items, "truncated": len(items) >= limit}


# ---------- Đọc file đính kèm ----------

TEXT_EXTS = {".txt", ".md", ".csv", ".tsv", ".json", ".log", ".xml", ".py", ".ini", ".yaml", ".yml"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
# Định dạng cũ / khác: chuyển sang dạng mới bằng Word / Excel / PowerPoint (hoặc WPS) trên máy rồi mới đọc
LEGACY = {".doc": "docx", ".rtf": "docx", ".odt": "docx", ".wps": "docx", ".dot": "docx", ".docm": "docx",
          ".ods": "xlsx", ".et": "xlsx", ".xlsb": "xlsx",
          ".ppt": "pptx", ".odp": "pptx", ".dps": "pptx", ".pps": "pptx", ".ppsx": "pptx"}
MAX_ROWS = 300


def decode_bytes(data: bytes) -> str:
    """Đoán bảng mã: BOM UTF-8/16, UTF-8, rồi bảng mã Windows tiếng Việt (cp1258, hay gặp ở CSV xuất từ Excel)."""
    import unicodedata
    if data[:4] in (b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff"):
        return data.decode("utf-32")
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    if data[:3] == b"\xef\xbb\xbf":
        return data[3:].decode("utf-8", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return unicodedata.normalize("NFC", data.decode("cp1258", errors="replace"))


def _cell(v) -> str:
    import datetime
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%d/%m/%Y") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, datetime.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _sheet_text(title, rows, n_rows, n_cols, path, tail=8) -> list[str]:
    """Hiện MAX_ROWS dòng đầu + vài dòng cuối (dòng Tổng cộng thường nằm cuối bảng)."""
    import collections

    def fmt(i, row):
        cells = list(row)
        while cells and cells[-1] == "":
            cells.pop()
        return f"{i + 1}\t" + "\t".join(cells) if i >= MAX_ROWS else "\t".join(cells)

    out = [f"## Sheet: {title} ({n_rows} dòng x {n_cols} cột)"]
    last = collections.deque(maxlen=tail)
    for i, row in enumerate(rows):
        if i < MAX_ROWS:
            out.append(fmt(i, row))
        else:
            last.append(fmt(i, row))
    if last:
        skipped = n_rows - MAX_ROWS - len(last)
        if skipped > 0:
            out.append(f"...(bỏ qua {skipped} dòng ở giữa — xử lý toàn bộ bằng pandas: "
                       f"pd.read_excel(r'{path}', sheet_name='{title}'))")
        out.append("(các dòng cuối, kèm số dòng:)")
        out += list(last)
    return out


def _xl_display(c) -> str:
    """Giá trị ô kèm định dạng có ý nghĩa: 0.15 + '0%' → '15%'; 123 + '00000' → '00123'."""
    v, fmt = c.value, (getattr(c, "number_format", "") or "")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if "%" in fmt:
            dec = len(fmt.split(".")[1].rstrip("%")) if "." in fmt else 0
            return f"{v * 100:.{dec}f}%"
        if re.fullmatch(r"0{2,}", fmt) and float(v).is_integer():
            return str(int(v)).zfill(len(fmt))
    return _cell(v)


def _read_xlsx(p: Path) -> str:
    import openpyxl
    vals = openpyxl.load_workbook(p, data_only=True, read_only=True)
    forms = openpyxl.load_workbook(p, data_only=False, read_only=True)
    parts = []
    for ws_v, ws_f in zip(vals.worksheets, forms.worksheets):
        rows = ([_xl_display(cv) if cv.value is not None
                 else (cf.value if isinstance(cf.value, str) and cf.value.startswith("=") else "")
                 for cv, cf in zip(rv, rf)]
                for rv, rf in zip(ws_v.iter_rows(), ws_f.iter_rows()))
        parts += _sheet_text(ws_v.title, rows, ws_v.max_row or 0, ws_v.max_column or 0, p)
    vals.close()
    forms.close()
    return "\n".join(parts) + "\n(Ô dạng =... là công thức chưa được Excel tính lại.)"


def _read_xls(p: Path) -> str:
    import xlrd
    book = xlrd.open_workbook(p)
    parts = []
    for sh in book.sheets():
        rows = ([_cell(xlrd.xldate_as_datetime(c.value, book.datemode)) if c.ctype == xlrd.XL_CELL_DATE
                 else _cell(c.value) for c in sh.row(r)] for r in range(sh.nrows))
        parts += _sheet_text(sh.name, rows, sh.nrows, sh.ncols, p)
    return "\n".join(parts)


def _read_docx(p: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(p)
    parts, seen = [], set()

    def cell_text(c):
        own = [x.text for x in c.paragraphs if x.text.strip()]
        own += [table_text(t) for t in c.tables]  # bảng lồng trong ô
        return " / ".join(own)

    def table_text(t):
        rows = []
        for row in t.rows:
            cells = []
            for c in row.cells:  # ô gộp bị python-docx trả lặp lại; bỏ trùng liền kề
                if not cells or cells[-1]._tc is not c._tc:
                    cells.append(c)
            rows.append("| " + " | ".join(cell_text(c) for c in cells) + " |")
        return "\n".join(rows)

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    def revisions(el) -> str:
        """Track Changes: phần đề xuất thêm / xoá kèm tác giả (python-docx bỏ qua phần này)."""
        out = []
        for kind, tag, label in (("ins", "t", "Đề xuất THÊM"), ("del", "delText", "Đề xuất XOÁ")):
            for r in el.iter(W + kind):
                txt = "".join(x.text or "" for x in r.iter(W + tag))
                if txt.strip():
                    out.append(f"[{label} — {r.get(W + 'author') or '?'}: «{txt}»]")
        return " ".join(out)

    for sec in d.sections:  # đầu / chân trang (thường, trang đầu, trang chẵn)
        for part, label in ((sec.header, "Đầu trang"), (sec.first_page_header, "Đầu trang (trang 1)"),
                            (sec.even_page_header, "Đầu trang (trang chẵn)"),
                            (sec.footer, "Chân trang"), (sec.first_page_footer, "Chân trang (trang 1)"),
                            (sec.even_page_footer, "Chân trang (trang chẵn)")):
            if part.is_linked_to_previous:
                continue
            txt = "\n".join(x.text for x in part.paragraphs if x.text.strip())
            if txt and txt not in seen:
                seen.add(txt)
                parts.append(f"[{label}] {txt}")
    for el in d.element.body.iterchildren():  # đúng thứ tự đoạn văn / bảng như trong file
        if el.tag.endswith("}p"):
            para = Paragraph(el, d)
            rev = revisions(el)
            if para.text.strip() or rev:
                style = para.style.name if para.style is not None else ""
                body = (para.text + (" " + rev if rev else "")).strip()
                if style.startswith("Heading") and style[-1:].isdigit():
                    parts.append("#" * int(style[-1]) + " " + body)
                else:
                    parts.append(("- " if "List" in style else "") + body)
        elif el.tag.endswith("}tbl"):
            parts.append(table_text(Table(el, d)))
            rev = revisions(el)
            if rev:
                parts.append(rev)
    if any("Đề xuất" in x for x in parts):
        parts.insert(0, "[Tài liệu có Track Changes: phần «Đề xuất THÊM/XOÁ» chưa được chấp nhận — đừng coi là nội dung chính thức.]")
    return "\n".join(parts)


def _read_pptx(p: Path) -> str:
    from pptx import Presentation

    def shape_text(sh):
        if sh.shape_type == 6:  # nhóm hình
            for x in sh.shapes:
                yield from shape_text(x)
        elif getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                yield "| " + " | ".join(c.text for c in row.cells) + " |"
        elif getattr(sh, "has_chart", False) and sh.has_chart:
            ch = sh.chart
            title = ch.chart_title.text_frame.text if ch.has_title and ch.chart_title.has_text_frame else ""
            yield f"[Biểu đồ{': ' + title if title else ''}]"
            for plot in ch.plots:
                cats = [str(c) for c in plot.categories]
                yield "Nhãn: " + ", ".join(cats)
                for ser in plot.series:
                    yield f"Chuỗi «{ser.name}»: " + ", ".join(_cell(v) for v in ser.values)
        elif sh.has_text_frame and sh.text_frame.text.strip():
            yield sh.text_frame.text

    parts = []
    for i, slide in enumerate(Presentation(p).slides, 1):
        parts.append(f"## Slide {i}")
        for sh in slide.shapes:
            parts += list(shape_text(sh))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append(f"[Ghi chú] {slide.notes_slide.notes_text_frame.text}")
    return "\n".join(parts)


def pdf_page_images(path: str, max_pages: int = 6, scale: float = 1.6, pages: list[int] | None = None) -> list[str]:
    """Dựng ảnh các trang PDF (để AI 'nhìn' bản scan hoặc kiểm tra bố cục). pages: số trang (bắt đầu từ 1)."""
    import pypdfium2 as pdfium
    out = IMAGES_DIR / "pages" / uuid.uuid4().hex[:8]
    out.mkdir(parents=True, exist_ok=True)
    pdf = pdfium.PdfDocument(path)
    wanted = [n for n in (pages or range(1, max_pages + 1)) if 1 <= n <= len(pdf)]
    paths = []
    for n in wanted:
        dest = out / f"trang_{n:02d}.jpg"
        pdf[n - 1].render(scale=scale).to_pil().convert("RGB").save(dest, quality=85)
        paths.append(str(dest))
    pdf.close()
    return paths


SCAN_MARK = "[Trang {n}: dạng scan/ảnh — không có lớp chữ]"
MAX_SCAN_PAGES = 10  # số trang scan gửi ảnh kèm mỗi lần; còn lại nói rõ để AI xem tiếp bằng preview_file


def _read_pdf(p: Path) -> str:
    from pypdf import PdfReader
    pages = [page.extract_text() or "" for page in PdfReader(p).pages]
    out = []
    for i, t in enumerate(pages, 1):
        out.append(f"--- Trang {i} ---")
        out.append(SCAN_MARK.format(n=i) if len(t.strip()) < 25 else t)
    return "\n".join(out)


def scanned_pages(text: str | None) -> list[int]:
    """Số các trang PDF dạng scan trong bản trích của extract_file_text."""
    return [int(n) for n in re.findall(r"\[Trang (\d+): dạng scan", text or "")]


def scan_attachments(path: str, text: str | None) -> tuple[list[str], str]:
    """Ảnh các trang scan (tối đa MAX_SCAN_PAGES) + ghi chú độ phủ chính xác. Không có trang scan → ([], "")."""
    pages = scanned_pages(text)
    if not pages:
        return [], ""
    sent = pages[:MAX_SCAN_PAGES]
    rest = pages[MAX_SCAN_PAGES:]
    note = f"[PDF có {len(pages)} trang dạng scan: {_ranges(pages)}. Đã gửi ảnh trang {_ranges(sent)}."
    if rest:
        note += (f" CHƯA gửi trang {_ranges(rest)} — xem tiếp bằng preview_file(path, pages=\"{_ranges(rest[:MAX_SCAN_PAGES])}\")"
                 f" trước khi trả lời về toàn bộ file.")
    return pdf_page_images(path, pages=sent), note + "]"


def _ranges(nums: list[int]) -> str:
    out, start = [], None
    for i, n in enumerate(nums):
        if start is None:
            start = n
        if i == len(nums) - 1 or nums[i + 1] != n + 1:
            out.append(str(start) if start == n else f"{start}-{n}")
            start = None
    return ",".join(out)


def extract_file_text(path: str, limit: int = 30000) -> str | None:
    p = Path(path)
    ext = p.suffix.lower()
    if ext in LEGACY:
        from office_helpers import convert_file
        p = Path(convert_file(str(p), LEGACY[ext], out_dir=str(OUTPUT_DIR.parent / "converted")))
        ext = p.suffix.lower()
    if ext == ".pdf":
        text = _read_pdf(p)
    elif ext == ".docx":
        text = _read_docx(p)
    elif ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        text = _read_xlsx(p)
    elif ext == ".xls":
        text = _read_xls(p)
    elif ext == ".pptx":
        text = _read_pptx(p)
    elif ext in (".html", ".htm"):
        soup = BeautifulSoup(decode_bytes(p.read_bytes()), "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        text = "\n".join(x.strip() for x in soup.get_text("\n").splitlines() if x.strip())
    elif ext in TEXT_EXTS:
        text = decode_bytes(p.read_bytes())
    else:
        return None
    return text[:limit] + ("\n...(đã cắt bớt — đọc toàn bộ bằng Python nếu cần)" if len(text) > limit else "")


def image_to_data_url(path: str, max_side: int = 1600) -> str:
    img = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
    img.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=85)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()
