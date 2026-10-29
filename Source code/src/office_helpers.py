"""Trợ giúp tạo Word (.docx), Excel (.xlsx), PDF đẹp, chuẩn + chuyển đổi định dạng. Được nạp sẵn vào sandbox."""
import contextlib
import os
import uuid
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as DocInches
from docx.shared import Pt as DocPt
from docx.shared import RGBColor as DocRGB
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR
from pptx_design_engine import FONT, THEMES, make_chart, output_path

__all__ = ["new_document", "add_title", "add_bullets", "add_table", "add_picture", "add_chart_to_doc",
           "save_document", "new_workbook", "write_table", "autofit_columns", "save_workbook",
           "convert_file", "recalculate_excel", "add_page_numbers", "add_toc", "docx_replace", "add_excel_chart",
           "register_vn_fonts"]

WORD_EXTS = {".doc", ".docx", ".docm", ".dot", ".dotx", ".rtf", ".odt", ".wps", ".txt", ".htm", ".html"}
EXCEL_EXTS = {".xls", ".xlsx", ".xlsm", ".xlsb", ".ods", ".et", ".csv"}
PPT_EXTS = {".ppt", ".pptx", ".pps", ".ppsx", ".odp", ".dps"}


# ---------- Chuyển đổi định dạng (dùng Word / Excel / PowerPoint hoặc WPS trên máy) ----------

def _office(progid):
    import win32com.client
    return win32com.client.dynamic.Dispatch(progid)


def _quit_if_idle(app, collection):
    with contextlib.suppress(Exception):
        if int(getattr(app, collection).Count) == 0:  # không tắt khi người dùng đang mở file khác
            app.Quit()


def convert_file(path, to="pdf", out_dir=None) -> str:
    """Chuyển định dạng: Word/Excel/PowerPoint → pdf; doc/rtf/odt → docx; xls/ods → xlsx; ppt → pptx; xlsx → csv...
    Mặc định lưu bản mới vào OUTPUT_DIR. Trả đường dẫn file mới."""
    import pythoncom
    src = Path(os.path.abspath(path))
    to = to.lower().lstrip(".")
    ext = src.suffix.lower()
    if out_dir:
        Path(out_dir).mkdir(parents=True, exist_ok=True)
        dest = Path(out_dir) / f"{src.stem}_{uuid.uuid4().hex[:6]}.{to}"
    else:
        dest = output_path(src.stem, "." + to)
    pythoncom.CoInitialize()  # code AI chạy ở luồng nền
    try:
        if ext in WORD_EXTS:
            app = _office("Word.Application")
            doc = app.Documents.Open(str(src))
            with contextlib.suppress(Exception):  # cập nhật mục lục, số trang... trước khi xuất
                for i in range(1, int(doc.TablesOfContents.Count) + 1):
                    doc.TablesOfContents(i).Update()
                doc.Fields.Update()
            if to == "pdf":
                doc.ExportAsFixedFormat(str(dest), 17)
            else:
                doc.SaveAs(str(dest), {"docx": 16, "doc": 0, "rtf": 6, "txt": 7, "html": 8, "odt": 23}[to])
            doc.Close(False)
            _quit_if_idle(app, "Documents")
        elif ext in EXCEL_EXTS:
            app = _office("Excel.Application")
            with contextlib.suppress(Exception):
                app.DisplayAlerts = False
            wb = app.Workbooks.Open(str(src))
            if to == "pdf":
                for sh in wb.Worksheets:  # in vừa khổ ngang 1 trang (chỉ khi xuất, không lưu vào file gốc)
                    with contextlib.suppress(Exception):
                        ps = sh.PageSetup
                        ps.Zoom = False
                        ps.FitToPagesWide = 1
                        ps.FitToPagesTall = False
                        if sh.UsedRange.Columns.Count > 6 or int(sh.ChartObjects().Count):
                            ps.Orientation = 2  # ngang
                wb.ExportAsFixedFormat(0, str(dest))
            else:
                wb.SaveAs(str(dest), {"xlsx": 51, "xls": 56, "csv": 62, "ods": 60}[to])
            wb.Close(False)
            _quit_if_idle(app, "Workbooks")
        elif ext in PPT_EXTS:
            app = _office("PowerPoint.Application")
            pres = app.Presentations.Open(str(src))
            pres.SaveAs(str(dest), {"pdf": 32, "pptx": 24, "ppt": 1}[to])
            pres.Close()
            _quit_if_idle(app, "Presentations")
        else:
            raise ValueError(f"Không chuyển được loại file {ext}")
    except KeyError:
        raise ValueError(f"Không hỗ trợ chuyển {ext} sang .{to}") from None
    finally:
        pythoncom.CoUninitialize()
    if not dest.exists():
        raise RuntimeError("Office trên máy không tạo được file (hãy kiểm tra Word/Excel/PowerPoint hoặc WPS)")
    print(f"Đã chuyển: {dest}")
    return str(dest)


def recalculate_excel(path) -> str:
    """Mở và lưu lại bằng Excel để mọi công thức có kết quả sẵn (openpyxl không tự tính công thức)."""
    import pythoncom
    pythoncom.CoInitialize()
    try:
        app = _office("Excel.Application")
        with contextlib.suppress(Exception):
            app.DisplayAlerts = False
        wb = app.Workbooks.Open(os.path.abspath(path))
        with contextlib.suppress(Exception):
            app.CalculateFull()
        wb.Save()
        wb.Close(False)
        _quit_if_idle(app, "Workbooks")
    finally:
        pythoncom.CoUninitialize()
    return str(path)


def _t(theme):
    return theme if isinstance(theme, dict) else THEMES.get(theme, THEMES["corporate_blue"])


# ---------- Word ----------

def _style_font(style, size=None, color=None, bold=None):
    f = style.font
    f.name = FONT
    if size:
        f.size = DocPt(size)
    if color:
        f.color.rgb = DocRGB.from_string(color)
    if bold is not None:
        f.bold = bold
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)  # font theme ghi đè font.name nên phải bỏ
    rfonts.set(qn("w:cs"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT)


def new_document(theme="corporate_blue"):
    """Document có sẵn Heading màu theo theme, chữ Segoe UI, giãn dòng 1.25, lề 2.2 cm."""
    t = _t(theme)
    doc = Document()
    doc._theme = t
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = DocInches(0.87)
        sec.top_margin = sec.bottom_margin = DocInches(0.8)
    normal = doc.styles["Normal"]
    _style_font(normal, 12, t["text"] if t is not THEMES["minimal_dark"] else "1F2937")
    pf = normal.paragraph_format
    pf.line_spacing = 1.25
    pf.space_after = DocPt(6)
    for name, size, color, before, after in (("Title", 26, t["strong"], 0, 4), ("Heading 1", 18, t["strong"], 18, 6),
                                             ("Heading 2", 14, t["title"] if t is not THEMES["minimal_dark"] else "374151", 12, 4),
                                             ("Heading 3", 12.5, t["strong"], 10, 3)):
        st = doc.styles[name]
        _style_font(st, size, color, True)
        st.paragraph_format.space_before = DocPt(before)
        st.paragraph_format.space_after = DocPt(after)
        st.paragraph_format.line_spacing = 1.1
    return doc


def add_title(doc, title, subtitle=""):
    doc.add_paragraph(title, style="Title")
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.runs[0].font.size = DocPt(13)
        p.runs[0].font.color.rgb = DocRGB.from_string(doc._theme["muted"])


def add_bullets(doc, items, numbered=False):
    for it in items:
        doc.add_paragraph(str(it), style="List Number" if numbered else "List Bullet")


def _shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def vn_number(v) -> str:
    """1234567 → '1.234.567'; 1234.56 → '1.234,56' (giữ đủ phần lẻ, không làm tròn)."""
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, float) and not v.is_integer():
        from decimal import Decimal
        ip, _, fp = format(Decimal(repr(v)), "f").partition(".")  # repr giữ đúng chữ số người dùng nhập
        txt = f"{int(ip):,}" + ("." + fp if fp else "")
        if ip.startswith("-0"):
            txt = "-" + txt
    else:
        txt = f"{int(v):,}"
    return txt.replace(",", "\0").replace(".", ",").replace("\0", ".")


def add_table(doc, headers, rows, col_widths_in=None, zebra=True):
    """Bảng có hàng tiêu đề tô màu đậm chữ trắng, hàng xen kẽ màu nhạt."""
    t = doc._theme
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.color.rgb = DocRGB(0xFF, 0xFF, 0xFF)
        _shade(cell, t["strong"])
    for r_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else vn_number(v)
            if isinstance(v, (int, float)):
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if zebra and r_i % 2:
                _shade(cells[i], t["soft"] if t is not THEMES["minimal_dark"] else "F3F4F6")
    if col_widths_in:
        for row in table.rows:
            for i, wd in enumerate(col_widths_in):
                row.cells[i].width = DocInches(wd)
    doc.add_paragraph()
    return table


def add_picture(doc, path, width_in=6.0, caption=None):
    doc.add_picture(str(path), width=DocInches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = doc.add_paragraph(caption, style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_chart_to_doc(doc, kind, labels, values, caption=None, unit="", width_in=6.0):
    # biểu đồ trên giấy trắng cần màu đậm
    t = THEMES["corporate_blue"] if doc._theme is THEMES["minimal_dark"] else doc._theme
    add_picture(doc, make_chart(kind, labels, values, t, w_in=width_in, h_in=width_in * 0.6, unit=unit), width_in, caption)


def save_document(doc, filename="tai_lieu.docx") -> str:
    path = output_path(filename, ".docx")
    doc.save(path)
    print(f"Đã lưu: {path}")
    return str(path)


# ---------- Excel ----------

def new_workbook():
    return openpyxl.Workbook()


def _col_index(headers, col):
    return headers.index(col) + 1 if isinstance(col, str) else int(col)


NO_TOTAL_WORDS = ("đơn giá", "giá bán", "giá nhập", "giá mua", "giá niêm yết", "unit price", "price", "tỷ lệ", "lãi suất")


def write_table(ws, title, headers, rows, money_cols=(), percent_cols=(), total_cols=None, total_label="Tổng cộng",
                theme="corporate_blue", subtitle=None, number_format="#,##0"):
    """Ghi bảng chuẩn: tiêu đề lớn gộp ô, header tô đậm chữ trắng, viền toàn bảng, định dạng số,
    hàng tổng dùng công thức SUM, tự căn độ rộng cột. Ô bắt đầu bằng "=" là công thức Excel.
    Trả về dict vị trí (header_row, first_row, last_row, total_row) để viết thêm công thức."""
    t = _t(theme)
    strong = t["strong"] if t is not THEMES["minimal_dark"] else "1F2937"
    n = len(headers)
    last_col = get_column_letter(n)
    thin = Side(style="thin", color="B8C2D1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    if n > 1:
        ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color=strong)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32
    hr = 3
    if subtitle:
        if n > 1:
            ws.merge_cells(f"A2:{last_col}2")
        ws["A2"] = subtitle
        ws["A2"].font = Font(name=FONT, size=11, italic=True, color=t["muted"])
        ws["A2"].alignment = Alignment(horizontal="center")
        hr = 4

    for i, h in enumerate(headers, 1):
        c = ws.cell(hr, i, h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=12)
        c.fill = PatternFill("solid", fgColor=strong)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[hr].height = 26

    money = {_col_index(headers, c) for c in money_cols}
    pct = {_col_index(headers, c) for c in percent_cols}
    soft = PatternFill("solid", fgColor=t["soft"] if t is not THEMES["minimal_dark"] else "F3F4F6")
    first = hr + 1
    for r_i, row in enumerate(rows):
        r = first + r_i
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v)
            c.font = Font(name=FONT, size=11)
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=isinstance(v, str) and len(v) > 40)
            if i in money:
                c.number_format = number_format
            elif i in pct:
                c.number_format = "0.0%"
            if r_i % 2:
                c.fill = soft
    last = first + len(rows) - 1

    total_row = None
    if total_cols is None:  # mặc định: cộng các cột tiền, trừ cột đơn giá / tỷ lệ (cộng lại vô nghĩa)
        total_cols = [c for c in money_cols if not any(w in str(headers[_col_index(headers, c) - 1]).lower()
                                                       for w in NO_TOTAL_WORDS)]
    sums = {_col_index(headers, c) for c in total_cols}
    if sums and rows:
        total_row = last + 1
        for i in range(1, n + 1):
            c = ws.cell(total_row, i)
            c.border = Border(left=thin, right=thin, bottom=thin, top=Side(style="medium", color=strong))
            c.font = Font(name=FONT, bold=True, size=12, color=strong)
            if i in sums:
                col = get_column_letter(i)
                c.value = f"=SUM({col}{first}:{col}{last})"
                c.number_format = number_format if i in money else "General"
        label_col = next((i for i in range(1, n + 1) if i not in sums), None)
        if label_col:  # không ghi nhãn đè lên ô tổng (khi cột số nằm ở cột A)
            ws.cell(total_row, label_col, total_label)

    ws.freeze_panes = ws.cell(hr + 1, 1)
    # Khi in / xuất PDF: vừa 1 trang ngang, lặp dòng tiêu đề bảng ở mỗi trang
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    if n > 5:
        ws.page_setup.orientation = "landscape"
    ws.print_title_rows = f"{hr}:{hr}"
    ws.print_options.horizontalCentered = True
    autofit_columns(ws)
    return {"header_row": hr, "first_row": first, "last_row": last, "total_row": total_row}


def autofit_columns(ws, min_width=8, max_width=60):
    """Căn độ rộng cột theo nội dung để không hiện ###."""
    merged = {f"{get_column_letter(col)}{row}" for rng in ws.merged_cells.ranges for row, col in rng.cells}
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None or c.coordinate in merged:
                continue
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                length = 14
            elif isinstance(v, (int, float)):
                length = len(f"{v:,.0f}") + 2
            else:
                length = max(len(line) for line in str(v).splitlines() or [""])
            widths[c.column_letter] = max(widths.get(c.column_letter, 0), length)
    for col, length in widths.items():
        ws.column_dimensions[col].width = max(min_width, min(max_width, length * 1.15 + 3))


def save_workbook(wb, filename="bang_tinh.xlsx") -> str:
    path = output_path(filename, ".xlsx")
    wb.save(path)
    print(f"Đã lưu: {path}")
    return str(path)


# ---------- Word nâng cao ----------

def _fld(paragraph, instr, placeholder=""):
    """Chèn trường tự động (PAGE, NUMPAGES, TOC...) vào đoạn văn."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = placeholder
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def add_page_numbers(doc, template="Trang {PAGE} / {NUMPAGES}", align="center"):
    """Số trang ở chân trang mọi section."""
    for sec in doc.sections:
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                       "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        for i, chunk in enumerate(template.replace("{NUMPAGES}", "\0N\0").replace("{PAGE}", "\0P\0").split("\0")):
            if chunk == "P":
                _fld(p, "PAGE", "1")
            elif chunk == "N":
                _fld(p, "NUMPAGES", "1")
            elif chunk:
                p.add_run(chunk)


def add_toc(doc, title="Mục lục", levels="1-3"):
    """Mục lục tự động theo Heading 1–3. Word/WPS hỏi cập nhật khi mở file (hoặc nhấn F9)."""
    head = doc.add_paragraph()  # không dùng Heading để chữ "Mục lục" không tự lọt vào mục lục
    run = head.add_run(title)
    h1 = doc.styles["Heading 1"].font
    run.bold, run.font.size, run.font.color.rgb = True, h1.size, h1.color.rgb
    _fld(doc.add_paragraph(), rf'TOC \o "{levels}" \h \z \u', "Mở file và chọn cập nhật để hiện mục lục.")
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)


def _iter_paragraphs(doc):
    def from_container(c):
        yield from c.paragraphs
        for t in c.tables:
            for row in t.rows:
                for cell in row.cells:
                    yield from from_container(cell)
    yield from from_container(doc)
    for sec in doc.sections:
        for part in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer):
            yield from from_container(part)


def docx_replace(doc, mapping: dict) -> int:
    """Thay chữ hàng loạt (thân, bảng, đầu/chân trang), giữ định dạng của run đầu. Dùng cho trộn thư:
    docx_replace(doc, {"{{ho_ten}}": "Nguyễn Văn An", ...}). Trả số chỗ đã thay."""
    count = 0
    for p in _iter_paragraphs(doc):
        for old, new in mapping.items():
            if old and old in p.text:
                count += _replace_in_runs(p, old, str(new))
    return count


def _replace_in_runs(p, old: str, new: str) -> int:
    """Placeholder có thể bị Word chia qua nhiều run: chữ thay vào run chứa ký tự đầu, chỉ cắt đúng phần
    placeholder ở các run sau, các run khác (vd. chữ đậm phía sau) giữ nguyên."""
    count, pos = 0, 0
    while True:
        runs = p.runs
        texts = [r.text for r in runs]
        full = "".join(texts)
        idx = full.find(old, pos)
        if idx < 0:
            return count
        end, start = idx + len(old), 0
        for r, t in zip(runs, texts):
            stop = start + len(t)
            if stop > idx and start < end:
                a, b = max(idx, start) - start, min(end, stop) - start
                r.text = t[:a] + (new if start <= idx < stop else "") + t[b:]
            start = stop
        count += 1
        pos = idx + len(new)


# ---------- Excel nâng cao ----------

def add_excel_chart(ws, kind="col", data_cols=("B",), cat_col="A", header_row=1, first_row=2, last_row=None,
                    anchor=None, title="", y_title="", width=18, height=9):
    """Biểu đồ Excel thật (sửa được trong Excel). kind: col (cột đứng) | bar (thanh ngang) | line | pie.
    data_cols: các cột số liệu (chữ cái hoặc số), cat_col: cột nhãn; dòng header_row là tên chuỗi."""
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.utils import column_index_from_string as ci
    idx = lambda c: c if isinstance(c, int) else ci(c)
    last_row = last_row or ws.max_row
    chart = {"bar": BarChart, "col": BarChart, "line": LineChart, "pie": PieChart}[kind]()
    if kind in ("bar", "col"):
        chart.type = "col" if kind == "col" else "bar"
        chart.grouping = "clustered"
    from openpyxl.chart import Series
    for c in data_cols:
        name = ws.cell(header_row, idx(c)).value
        chart.series.append(Series(Reference(ws, min_col=idx(c), min_row=first_row, max_row=last_row),
                                   title=str(name) if name is not None else None))
    chart.set_categories(Reference(ws, min_col=idx(cat_col), min_row=first_row, max_row=last_row))
    chart.title = title or None
    if kind != "pie" and y_title:
        chart.y_axis.title = y_title
    chart.width, chart.height = width, height
    ws.add_chart(chart, anchor or f"{get_column_letter(ws.max_column + 2)}{header_row}")
    return chart


# ---------- PDF trực tiếp (khi máy không có Office) ----------

def register_vn_fonts() -> str:
    """Đăng ký font Windows có đủ dấu tiếng Việt cho reportlab. Trả tên font thường (bản đậm: tên + '-Bold')."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fonts = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    for name, regular, bold in (("Arial", "arial.ttf", "arialbd.ttf"), ("SegoeUI", "segoeui.ttf", "segoeuib.ttf")):
        if (fonts / regular).exists():
            pdfmetrics.registerFont(TTFont(name, str(fonts / regular)))
            pdfmetrics.registerFont(TTFont(name + "-Bold", str(fonts / (bold if (fonts / bold).exists() else regular))))
            return name
    raise RuntimeError(r"Không tìm thấy font tiếng Việt trong Windows\Fonts")
