"""Bộ máy thiết kế PowerPoint (pptx_helper): theme, nhiều layout cao cấp, icon, hiệu ứng, chống tràn chữ.

Được nạp sẵn vào sandbox `execute_python_file_processor`. API chính là lớp `Deck`.
"""
import copy
import math
import uuid
from pathlib import Path

from lxml import etree
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from config import IMAGES_DIR, OUTPUT_DIR

__all__ = ["THEMES", "FONT", "TITLE_FONT", "ICONS", "Deck", "add_text", "add_box", "add_image", "add_icon",
           "crop_to_fill", "fit_font_size", "make_chart", "output_path", "set_alpha", "add_shadow", "gradient_fill",
           "make_theme"]

FONT = "Segoe UI"
TITLE_FONT = "Segoe UI Semibold"
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.7
CONTENT_Y, CONTENT_H = 1.85, 4.95

THEMES = {
    "navy_gold": dict(name="Modern Navy & Gold", bg="F8F6F1", card="FFFFFF", border="E7E1D4", title="14213D",
                      text="1F2937", muted="5B6475", accent="D9A21B", dark="0F1A33", on_dark="FFFFFF",
                      on_dark_muted="C5CEE0", strong="14213D", soft="F2EBDA",
                      chart=["14213D", "D9A21B", "4F6D9A", "B7791F", "9AAAC4"]),
    "emerald_cream": dict(name="Emerald & Warm Cream", bg="FBF7EF", card="FFFFFF", border="E8DFCC", title="0B4A2E",
                          text="1F2A24", muted="56635B", accent="C98A12", dark="0A3D26", on_dark="FFFFFF",
                          on_dark_muted="C9E2D3", strong="0F5132", soft="E9F3EC",
                          chart=["0F5132", "C98A12", "3F8F65", "A86B12", "8DBFA2"]),
    "corporate_blue": dict(name="Corporate Blue & Slate", bg="F5F7FB", card="FFFFFF", border="DCE3EE", title="13265C",
                           text="0F172A", muted="475569", accent="0EA5E9", dark="0B1B45", on_dark="FFFFFF",
                           on_dark_muted="C3CDEB", strong="1D4ED8", soft="E6EEFD",
                           chart=["1D4ED8", "0EA5E9", "64748B", "13265C", "93C5FD"]),
    "sunset_terracotta": dict(name="Sunset Terracotta", bg="FDF6F0", card="FFFFFF", border="F0DCCB", title="6B240C",
                              text="2B1A12", muted="6B5446", accent="E8920B", dark="4A1A09", on_dark="FFFFFF",
                              on_dark_muted="F2CDB4", strong="C2410C", soft="FBE7DA",
                              chart=["C2410C", "E8920B", "7C2D12", "E07A5F", "F2B880"]),
    "medical_teal": dict(name="Medical Clean Teal", bg="F3FAF9", card="FFFFFF", border="D3E9E6", title="0B4141",
                         text="102A2A", muted="4A6363", accent="10A89A", dark="073B3B", on_dark="FFFFFF",
                         on_dark_muted="B9E3DE", strong="0F766E", soft="DFF3F0",
                         chart=["0F766E", "10A89A", "0B4141", "5EAAA8", "99D5CF"]),
    "minimal_dark": dict(name="Minimalist Dark", bg="111318", card="1A1E26", border="2A303C", title="F5F6F8",
                         text="E6E8EC", muted="A3AAB8", accent="7DD3FC", dark="0A0B0E", on_dark="F5F6F8",
                         on_dark_muted="A3AAB8", strong="38BDF8", soft="1A1E26",
                         chart=["7DD3FC", "F5F6F8", "64748B", "38BDF8", "CBD5E1"]),
}

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_PARA_GAP = 0.35  # khoảng cách giữa đoạn, tính theo cỡ chữ
WARNINGS: list[str] = []  # chữ bị thu quá nhỏ / bị cắt; Deck.save() in ra để AI rút gọn nội dung


def _icons():
    import icons
    return icons


ICONS = _icons().SLIDE_ICONS


def _theme(theme):
    return theme if isinstance(theme, dict) else THEMES.get(theme, THEMES["corporate_blue"])


def _mix(a: str, b: str, t: float) -> str:
    """Trộn 2 màu hex: t=0 → a, t=1 → b."""
    a, b = a.lstrip("#"), b.lstrip("#")
    return "".join(f"{round(int(a[i:i + 2], 16) * (1 - t) + int(b[i:i + 2], 16) * t):02X}" for i in (0, 2, 4))


def make_theme(primary, accent, bg="#F8FAFC", fg="#1E293B", card="#FFFFFF", dark=None, name="Custom"):
    """Tạo theme đầy đủ từ vài màu gốc (ví dụ bảng màu thương hiệu)."""
    p, a, bg, fg, card = (c.lstrip("#").upper() for c in (primary, accent, bg, fg, card))
    dark = (dark or _mix(p, "000000", 0.55)).lstrip("#").upper()
    return dict(name=name, bg=bg, card=card, border=_mix(bg, fg, 0.12), title=_mix(p, "000000", 0.35), text=fg,
                muted=_mix(fg, bg, 0.38), accent=a, dark=dark, on_dark="FFFFFF", on_dark_muted=_mix("FFFFFF", dark, 0.25),
                strong=p, soft=_mix(p, "FFFFFF", 0.88),
                chart=[p, a, _mix(p, "FFFFFF", 0.45), _mix(a, "000000", 0.25), _mix(fg, "FFFFFF", 0.5)])


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def output_path(filename: str, ext: str) -> Path:
    """Đường dẫn trong OUTPUT_DIR, tự thêm số nếu trùng tên."""
    stem = Path(filename).stem or "tai_lieu"
    path = OUTPUT_DIR / f"{stem}{ext}"
    i = 2
    while path.exists():
        path = OUTPUT_DIR / f"{stem} ({i}){ext}"
        i += 1
    return path


# ---------- Hiệu ứng (chỉnh XML trực tiếp vì python-pptx chưa hỗ trợ) ----------

def set_alpha(shape, opacity: float):
    """Độ trong suốt cho màu nền đặc của shape (opacity 0..1)."""
    clr = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    for old in clr.findall(qn("a:alpha")):
        clr.remove(old)
    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(opacity * 100000)))
    return shape


def gradient_fill(shape, color, stops=((0, 0.92), (1, 0.15)), angle=0):
    """Phủ chuyển màu một màu với độ trong suốt khác nhau: stops = ((vị_trí 0..1, độ_đục 0..1), ...)."""
    shape.fill.gradient()
    shape.fill.gradient_angle = angle
    gs_lst = shape.fill._fill._gradFill.gsLst
    template = gs_lst[0]
    for gs in list(gs_lst):
        gs_lst.remove(gs)
    for pos, opacity in stops:
        gs = copy.deepcopy(template)
        for child in list(gs):
            gs.remove(child)
        gs.set("pos", str(int(pos * 100000)))
        clr = etree.SubElement(gs, qn("a:srgbClr"))
        clr.set("val", color)
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(opacity * 100000)))
        gs_lst.append(gs)
    shape.line.fill.background()
    return shape


def add_shadow(shape, blur=0.22, dist=0.05, opacity=0.14):
    """Bóng đổ mềm phía dưới (thẻ nổi nhẹ khỏi nền)."""
    sp_pr = shape._element.spPr
    for old in sp_pr.findall(qn("a:effectLst")):
        sp_pr.remove(old)
    eff = etree.SubElement(sp_pr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(blur * 914400)), dist=str(int(dist * 914400)),
                          dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
    etree.SubElement(clr, qn("a:alpha"), val=str(int(opacity * 100000)))
    return shape


def _round_picture(pic, radius=0.06):
    geom = pic._element.spPr.find(qn("a:prstGeom"))
    geom.set("prst", "roundRect")
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    for g in list(av):
        av.remove(g)
    etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {int(radius * 100000)}")


# ---------- Chống tràn chữ ----------

def _fits(text: str, w_in: float, h_in: float, pt: float, line_spacing: float, bold: bool) -> bool:
    # shortcut: ước lượng bề rộng ký tự trung bình của Segoe UI (~0.52em, đậm ~0.57em), không đo glyph thật
    chars_per_line = max(1, int(w_in * 72 / (pt * (0.57 if bold else 0.52))))
    paras = text.split("\n")
    lines = sum(max(1, math.ceil(len(p) * 1.08 / chars_per_line)) for p in paras)
    # 1.3: PowerPoint/WPS dàn dòng tiếng Việt (có dấu) cao hơn ~1.2em danh nghĩa — đo trên ảnh xuất thật
    need = (lines * line_spacing * 1.3 + (len(paras) - 1) * _PARA_GAP) * pt / 72
    return need <= h_in


def fit_font_size(text, w_in, h_in, max_pt=24, min_pt=11, line_spacing=1.15, bold=False) -> int:
    for pt in range(int(max_pt), int(min_pt) - 1, -1):
        if _fits(text, w_in, h_in, pt, line_spacing, bold):
            return pt
    return int(min_pt)


def _truncate(text, w_in, h_in, pt, line_spacing, bold):
    if _fits(text, w_in, h_in, pt, line_spacing, bold):
        return text
    words = text.split(" ")
    while len(words) > 3 and not _fits(" ".join(words) + "…", w_in, h_in, pt, line_spacing, bold):
        words.pop()
    return " ".join(words).rstrip(",.;:") + "…"


def add_text(slide, text, x, y, w, h, size=20, min_size=11, color="1F2937", bold=False,
             align="left", anchor="top", line_spacing=1.15, bullets=False, font=FONT, spacing=None):
    """Hộp chữ tự giảm cỡ chữ (và cắt bớt khi vẫn quá dài) để không bao giờ tràn khung."""
    if isinstance(text, (list, tuple)):
        text = "\n".join(("•  " if bullets else "") + str(t) for t in text)
    text = str(text).strip()
    pad = 0.06
    iw, ih = w - 2 * pad, h - 2 * pad
    heavy = bold or font == TITLE_FONT
    pt = fit_font_size(text, iw, ih, size, min_size, line_spacing, heavy)
    short = _truncate(text, iw, ih, pt, line_spacing, heavy)
    if short != text:
        WARNINGS.append(f"chữ quá dài, đã bị cắt: “{text[:60]}…”")
    elif pt <= min_size and size > min_size + 4:
        WARNINGS.append(f"chữ phải thu nhỏ còn {pt}pt: “{text[:60]}…”")
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(pad)
    tf.vertical_anchor = _ANCHOR[anchor]
    for i, line in enumerate(short.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[align]
        p.line_spacing = line_spacing
        if i:
            p.space_before = Pt(pt * _PARA_GAP)
        run = p.add_run()
        run.text = line
        f = run.font
        f.name, f.size, f.bold = font, Pt(pt), bold
        f.color.rgb = rgb(color)
        if spacing is not None:  # giãn chữ (đơn vị 1/100 pt), dùng cho nhãn viết hoa
            run._r.get_or_add_rPr().set("spc", str(int(spacing)))
    return box


def add_box(slide, x, y, w, h, fill, line=None, rounded=True, radius=0.06, shape=None, alpha=None, shadow=False):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    shp.shadow.inherit = False
    if alpha is not None:
        set_alpha(shp, alpha)
    if shadow:
        add_shadow(shp)
    return shp


def add_icon(slide, name, x, y, size, color, badge=None, badge_alpha=1.0):
    """Icon nét mảnh (xem ICONS); badge = màu vòng tròn nền phía sau."""
    if badge:
        add_box(slide, x, y, size, size, badge, shape=MSO_SHAPE.OVAL, alpha=badge_alpha)
    if name not in _icons().PATHS:
        return None
    out = IMAGES_DIR / "icons"
    out.mkdir(exist_ok=True)
    path = out / f"{name}_{color}.png"
    if not path.exists():
        _icons().render_png(name, "#" + color, str(path))
    inset = size * (0.26 if badge else 0)
    return slide.shapes.add_picture(str(path), Inches(x + inset), Inches(y + inset),
                                    Inches(size - 2 * inset), Inches(size - 2 * inset))


# ---------- Ảnh ----------

def crop_to_fill(path, w_in, h_in) -> str:
    """Cắt ảnh (Pillow) vừa khít tỉ lệ khung, không méo hình."""
    img = ImageOps.exif_transpose(Image.open(path))
    if img.mode in ("RGBA", "LA", "P"):
        bg = Image.new("RGB", img.size, "white")
        rgba = img.convert("RGBA")
        bg.paste(rgba, mask=rgba.split()[3])
        img = bg
    img = img.convert("RGB")
    target, (iw, ih) = w_in / h_in, img.size
    if iw / ih > target:
        nw = int(ih * target)
        left = (iw - nw) // 2
        img = img.crop((left, 0, left + nw, ih))
    else:
        nh = int(iw / target)
        top = int((ih - nh) * 0.35)  # hơi lệch lên trên để giữ phần đầu/chủ thể
        img = img.crop((0, top, iw, top + nh))
    # giới hạn kích thước để file pptx nhẹ
    max_px = int(max(w_in, h_in) * 220)
    if max(img.size) > max_px:
        img.thumbnail((max_px, max_px))
    out = IMAGES_DIR / "crops"
    out.mkdir(exist_ok=True)
    path = out / f"crop_{uuid.uuid4().hex[:8]}.jpg"
    img.save(path, quality=90)
    return str(path)


def add_image(slide, path, x, y, w, h, theme="corporate_blue", radius=0.0):
    """Chèn ảnh đã cắt vừa khung (radius > 0 thì bo góc); thiếu ảnh thì vẽ khung trang trí thay thế."""
    if path and Path(path).is_file():
        try:
            pic = slide.shapes.add_picture(crop_to_fill(path, w, h), Inches(x), Inches(y), Inches(w), Inches(h))
            if radius:
                _round_picture(pic, radius)
            return pic
        except Exception:
            pass
    t = _theme(theme)
    add_box(slide, x, y, w, h, t["soft"], rounded=bool(radius), radius=radius or 0.06)
    return add_icon(slide, "image", x + w / 2 - 0.5, y + h / 2 - 0.5, 1.0, t["muted"])


# ---------- Biểu đồ ----------

def make_chart(kind, labels, values, theme="corporate_blue", w_in=7.6, h_in=5.0, unit="", title=None) -> str:
    """Vẽ biểu đồ matplotlib theo màu theme, PNG 300 DPI nền trong suốt. kind: bar | barh | line | pie | donut."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    t = _theme(theme)
    c = lambda h: "#" + h
    plt.rcParams.update({"font.family": ["Segoe UI", "Arial", "DejaVu Sans"], "font.size": 13})
    series = values if isinstance(values, dict) else {"": list(values)}
    fig, ax = plt.subplots(figsize=(w_in, h_in), dpi=300)
    fig.patch.set_alpha(0)
    ax.set_facecolor("none")
    colors = [c(x) for x in t["chart"]]
    fmt = lambda v: f"{v:,.0f}{unit}".replace(",", ".") if abs(v) >= 1000 else f"{v:g}{unit}"

    if kind in ("pie", "donut"):
        vals = list(series.values())[0]
        _, texts, autotexts = ax.pie(vals, labels=labels, colors=colors[:len(vals)], startangle=90, counterclock=False,
                                     autopct="%1.0f%%", pctdistance=0.78,
                                     wedgeprops=dict(width=0.42, edgecolor=c(t["bg"]), linewidth=3))
        for tx in texts:
            tx.set_color(c(t["text"]))
        for tx in autotexts:
            tx.set_color("white")
            tx.set_fontweight("bold")
        ax.axis("equal")
    else:
        n = len(series)
        idx = list(range(len(labels)))
        width = 0.7 / n
        for i, (name, vals) in enumerate(series.items()):
            col = colors[i % len(colors)]
            if kind == "line":
                ax.plot(labels, vals, color=col, linewidth=3.2, marker="o", markersize=8, label=name or None,
                        markerfacecolor=c(t["bg"]), markeredgewidth=2.5)
                ax.fill_between(labels, vals, min(vals), color=col, alpha=0.08)
                for xi, v in zip(labels, vals):
                    ax.annotate(fmt(v), (xi, v), textcoords="offset points", xytext=(0, 11), ha="center",
                                color=c(t["text"]), fontsize=11, fontweight="bold")
            else:
                pos = [p + (i - (n - 1) / 2) * width for p in idx]
                draw = ax.barh if kind == "barh" else ax.bar
                # nổi bật cột lớn nhất, các cột khác nhạt hơn (khi chỉ có 1 chuỗi)
                bar_colors = col if n > 1 else [col if v == max(vals) else col + "88" for v in vals]
                bars = draw(pos, vals, width * 0.9, color=bar_colors, label=name or None)
                ax.bar_label(bars, labels=[fmt(v) for v in vals], padding=5, color=c(t["text"]), fontsize=11,
                             fontweight="bold")
        if kind != "line":
            (ax.set_yticks if kind == "barh" else ax.set_xticks)(idx, labels)
        for side in ("top", "right", "left"):
            ax.spines[side].set_visible(False)
        ax.spines["bottom"].set_color(c(t["border"]))
        ax.tick_params(colors=c(t["muted"]), length=0)
        ax.grid(axis="x" if kind == "barh" else "y", color=c(t["border"]), linewidth=0.8)
        ax.set_axisbelow(True)
        if n > 1:
            ax.legend(frameon=False, labelcolor=c(t["text"]))
    if title:
        ax.set_title(title, color=c(t["title"]), fontweight="bold")
    fig.tight_layout()
    out = IMAGES_DIR / "charts"
    out.mkdir(exist_ok=True)
    path = out / f"chart_{uuid.uuid4().hex[:8]}.png"
    fig.savefig(path, dpi=300, transparent=True)
    plt.close(fig)
    return str(path)


# ---------- Deck ----------

def _item(item):
    """Chuẩn hoá 1 ý thành (icon, tiêu đề, nội dung). Nhận str, (tiêu đề, nội dung), (icon, tiêu đề, nội dung) hoặc dict."""
    if isinstance(item, dict):
        return item.get("icon"), str(item.get("title", "")), str(item.get("text", item.get("body", "")))
    if isinstance(item, (list, tuple)):
        if len(item) >= 3:
            return item[0], str(item[1]), str(item[2])
        return None, str(item[0]), str(item[1]) if len(item) > 1 else ""
    return None, "", str(item)


class Deck:
    """Bài thuyết trình 16:9 có theme đồng bộ, số trang, nhiều kiểu layout và hiệu ứng tinh tế."""

    def __init__(self, theme="corporate_blue", footer=""):
        self.t = _theme(theme)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
        self.footer = footer
        self.n = 0
        self.is_dark = self.t is THEMES["minimal_dark"]
        WARNINGS.clear()

    # nền tảng chung
    def _new(self, dark=False):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(self.t["dark"] if dark else self.t["bg"])
        return s

    def _decor(self, s, dark=False):
        """Hoạ tiết vòng tròn mờ ở góc — tạo chiều sâu mà không rối mắt."""
        t = self.t
        add_box(s, SLIDE_W - 3.2, -2.2, 5.2, 5.2, t["accent"], shape=MSO_SHAPE.OVAL, alpha=0.10 if dark else 0.07)
        add_box(s, SLIDE_W - 1.6, 4.9, 3.4, 3.4, t["on_dark"] if dark else t["strong"], shape=MSO_SHAPE.OVAL,
                alpha=0.05 if dark else 0.04)

    def _footer(self, s, dark=False):
        t = self.t
        col = t["on_dark_muted"] if dark else t["muted"]
        if not dark:
            add_box(s, MARGIN, SLIDE_H - 0.62, SLIDE_W - 2 * MARGIN, 0.012, t["border"], rounded=False)
        if self.footer:
            add_text(s, self.footer, MARGIN - 0.06, SLIDE_H - 0.58, 7, 0.4, size=11, color=col, anchor="middle")
        add_text(s, f"{self.n:02d}", SLIDE_W - MARGIN - 0.8, SLIDE_H - 0.58, 0.86, 0.4, size=11, color=col,
                 align="right", anchor="middle", font=TITLE_FONT)

    def _chrome(self, s, title, kicker=None):
        t = self.t
        self._decor(s)
        add_box(s, MARGIN, 0.62, 0.42, 0.07, t["accent"], rounded=True, radius=0.5)
        add_text(s, title, MARGIN - 0.06, 0.78, SLIDE_W - 2 * MARGIN - 1.5, 0.95, size=32, min_size=20,
                 color=t["title"], font=TITLE_FONT, anchor="middle")
        self._footer(s)

    def _card(self, s, x, y, w, h, icon, head, body, body_size=16, shadow=True):
        t = self.t
        add_box(s, x, y, w, h, t["card"], line=t["border"], radius=0.07 if h > 2 else 0.12, shadow=shadow)
        compact = h < 2.0
        pad = 0.22 if compact else 0.32
        top = y + pad
        if icon and not compact and (h < 2.9 or not head):
            # thẻ thấp: icon nằm cùng hàng tiêu đề để chừa chỗ cho nội dung
            add_icon(s, icon, x + pad, top, 0.56, t["strong"] if not self.is_dark else t["accent"],
                     badge=t["soft"] if not self.is_dark else t["border"])
            if head:
                add_text(s, head, x + pad + 0.7, top, w - 2 * pad - 0.64, 0.56, size=22, min_size=14,
                         color=t["title"], font=TITLE_FONT, anchor="middle")
            top += 0.74
            add_text(s, body, x + pad - 0.06, top, w - 2 * pad + 0.12, y + h - pad - top, size=body_size, min_size=11,
                     color=t["muted"] if not self.is_dark else t["text"], line_spacing=1.22)
            return
        if icon and not compact:
            add_icon(s, icon, x + pad, top, 0.62, t["strong"] if not self.is_dark else t["accent"],
                     badge=t["soft"] if not self.is_dark else t["border"])
            top += 0.82
        elif icon and compact:
            add_icon(s, icon, x + pad, y + h / 2 - 0.27, 0.54, t["strong"] if not self.is_dark else t["accent"],
                     badge=t["soft"] if not self.is_dark else t["border"])
            x, w = x + 0.66, w - 0.66
        if head:
            hh = 0.52 if compact else max(0.55, min(0.9, (y + h - top) * 0.32))
            add_text(s, head, x + pad - 0.06, top, w - 2 * pad + 0.12, hh, size=20 if compact else 22, min_size=14,
                     color=t["title"], font=TITLE_FONT)
            top += hh + (0 if compact else 0.04)
        add_text(s, body, x + pad - 0.06, top, w - 2 * pad + 0.12, y + h - pad - top, size=body_size, min_size=11,
                 color=t["muted"] if not self.is_dark else t["text"], line_spacing=1.22)

    def blank_slide(self, title):
        """Slide trống có sẵn tiêu đề + số trang. Trả (slide, (x, y, w, h)) vùng nội dung tính theo inch."""
        s = self._new()
        self._chrome(s, title)
        return s, (MARGIN, CONTENT_Y, SLIDE_W - 2 * MARGIN, CONTENT_H)

    # ===== các layout =====
    def hero(self, title, subtitle="", image=None, note="", style="full"):
        """Bìa. style="full": ảnh tràn slide + lớp phủ chuyển màu; "split": ảnh nửa phải."""
        t = self.t
        s = self._new(dark=True)
        if image and style == "full":
            add_image(s, image, 0, 0, SLIDE_W, SLIDE_H, t)
            gradient_fill(add_box(s, 0, 0, SLIDE_W, SLIDE_H, t["dark"], rounded=False), t["dark"],
                          stops=((0, 0.94), (0.55, 0.72), (1, 0.12)))
            tw = 7.4
        elif image:
            add_image(s, image, 6.9, 0.55, SLIDE_W - 6.9 - 0.55, SLIDE_H - 1.1, t, radius=0.05)
            tw = 5.8
        else:
            self._decor(s, dark=True)
            add_box(s, 8.6, 1.3, 4.9, 4.9, t["accent"], shape=MSO_SHAPE.OVAL, alpha=0.16)
            add_box(s, 9.6, 2.3, 2.9, 2.9, t["accent"], shape=MSO_SHAPE.OVAL, alpha=0.22)
            tw = 8.2
        add_box(s, 0.85, 1.55, 0.5, 0.07, t["accent"], radius=0.5)
        add_text(s, title, 0.8, 1.75, tw, 2.9, size=50, min_size=30, color=t["on_dark"], font=TITLE_FONT,
                 anchor="bottom", line_spacing=1.02)
        if subtitle:
            add_text(s, subtitle, 0.8, 4.8, tw, 1.3, size=21, min_size=14, color=t["on_dark_muted"], line_spacing=1.25)
        if note:
            add_text(s, note, 0.8, 6.5, tw, 0.45, size=13, color=t["on_dark_muted"])
        return s

    def section(self, title, subtitle="", number=None, image=None):
        """Slide chia chương: số chương lớn + tiêu đề, có thể kèm ảnh bên phải."""
        t = self.t
        s = self._new(dark=True)
        tw = 6.4 if image else 10
        if image:
            add_image(s, image, 7.4, 0, SLIDE_W - 7.4, SLIDE_H, t)
            gradient_fill(add_box(s, 7.4, 0, 2.2, SLIDE_H, t["dark"], rounded=False), t["dark"],
                          stops=((0, 1), (1, 0)))
        else:
            self._decor(s, dark=True)
        num = f"{number:02d}" if isinstance(number, int) else (number or f"{self.n:02d}")
        add_text(s, num, 0.8, 1.3, 4, 1.7, size=96, min_size=60, color=t["accent"], font=TITLE_FONT, anchor="bottom")
        add_text(s, title, 0.8, 3.15, tw, 1.6, size=42, min_size=26, color=t["on_dark"], font=TITLE_FONT)
        if subtitle:
            add_text(s, subtitle, 0.8, 4.8, tw, 1.2, size=20, min_size=14, color=t["on_dark_muted"])
        return s

    def agenda(self, title, items):
        """Mục lục: cột trái nền đậm chứa tiêu đề, cột phải danh sách đánh số."""
        t = self.t
        s = self._new()
        add_box(s, 0, 0, 4.6, SLIDE_H, t["dark"], rounded=False)
        add_box(s, 0.8, 1.6, 0.5, 0.07, t["accent"], radius=0.5)
        add_text(s, title, 0.75, 1.8, 3.4, 3.2, size=40, min_size=24, color=t["on_dark"], font=TITLE_FONT)
        rows = [_item(i) for i in items][:6]
        n = len(rows)
        top, avail = 0.9, SLIDE_H - 1.8
        rh = min(1.3, avail / max(n, 1))
        y0 = top + (avail - rh * n) / 2
        for i, (_, head, body) in enumerate(rows):
            y = y0 + i * rh
            add_text(s, f"{i + 1:02d}", 5.4, y, 1.0, rh, size=30, color=t["accent"], font=TITLE_FONT, anchor="middle")
            label, desc = (head, body) if head else (body, "")
            if desc:
                add_text(s, label, 6.5, y + 0.08, 6.1, rh * 0.48, size=25, min_size=15, color=t["title"],
                         font=TITLE_FONT, anchor="bottom")
                add_text(s, desc, 6.5, y + rh * 0.52, 6.1, rh * 0.44, size=17, min_size=12, color=t["muted"])
            else:
                add_text(s, label, 6.5, y, 6.1, rh, size=22, min_size=14, color=t["title"], font=TITLE_FONT,
                         anchor="middle")
            if i < n - 1:
                add_box(s, 5.4, y + rh - 0.01, 7.2, 0.012, t["border"], rounded=False)
        self._footer(s)
        return s

    def split(self, title, points, image=None, image_side="right"):
        """2 cột: các thẻ nội dung (có thể kèm icon) + ảnh bo góc."""
        t = self.t
        s, (x, y, w, h) = self.blank_slide(title)
        img_w = 5.3
        text_w = w - img_w - 0.45
        tx = x if image_side == "right" else x + img_w + 0.45
        ix = x + text_w + 0.45 if image_side == "right" else x
        add_image(s, image, ix, y, img_w, h, t, radius=0.05)
        items = [_item(p) for p in points][:4]
        if all(not head for _, head, _ in items):
            add_box(s, tx, y, text_w, h, t["card"], line=t["border"], radius=0.05, shadow=True)
            add_text(s, [b for _, _, b in items], tx + 0.32, y + 0.3, text_w - 0.64, h - 0.6, size=21, min_size=13,
                     color=t["text"], bullets=True, line_spacing=1.3)
        else:
            gap = 0.2
            ch = (h - gap * (len(items) - 1)) / len(items)
            for i, (icon, head, body) in enumerate(items):
                self._card(s, tx, y + i * (ch + gap), text_w, ch, icon, head, body, body_size=18)
        return s

    def cards(self, title, cards, columns=None):
        """Thẻ bo góc có bóng mềm: list[(icon, tiêu đề, nội dung)] hoặc list[(tiêu đề, nội dung)], 2–6 thẻ."""
        s, (x, y, w, h) = self.blank_slide(title)
        items = [_item(c) for c in cards][:6]
        n = len(items)
        cols = columns or (2 if n in (2, 4) else 3)
        rows = math.ceil(n / cols)
        gap = 0.3
        cw = (w - gap * (cols - 1)) / cols
        ch = (h - gap * (rows - 1)) / rows
        for i, (icon, head, body) in enumerate(items):
            r, c = divmod(i, cols)
            self._card(s, x + c * (cw + gap), y + r * (ch + gap), cw, ch, icon, head, body,
                       body_size=20 if rows == 1 else 18)
        return s

    def stats(self, title, items, note=""):
        """Số liệu lớn nổi bật: items = list[(giá_trị, nhãn, mô_tả)] 2–4 mục, ví dụ ("70%", "Người cao tuổi", "...")."""
        t = self.t
        s, (x, y, w, h) = self.blank_slide(title)
        rows = [tuple(i) + ("",) * (3 - len(i)) for i in items][:4]
        n = len(rows)
        gap = 0.35
        cw = (w - gap * (n - 1)) / n
        ch = min(h - (0.9 if note else 0.2), 3.9)
        y = y + (h - (0.9 if note else 0) - ch) / 2 - 0.1
        for i, (value, label, desc) in enumerate(rows):
            cx = x + i * (cw + gap)
            add_box(s, cx, y + 0.1, cw, ch, t["card"], line=t["border"], radius=0.06, shadow=True)
            add_box(s, cx + 0.35, y + 0.42, 0.5, 0.07, t["accent"], radius=0.5)
            add_text(s, str(value), cx + 0.28, y + 0.6, cw - 0.5, 1.5, size=66 if n <= 3 else 54, min_size=30,
                     color=t["strong"] if not self.is_dark else t["accent"], font=TITLE_FONT, anchor="middle")
            add_text(s, label, cx + 0.3, y + 2.15, cw - 0.6, 0.7, size=23, min_size=14, color=t["title"], font=TITLE_FONT)
            if desc:
                add_text(s, desc, cx + 0.3, y + 2.85, cw - 0.6, ch - 2.95, size=18, min_size=12, color=t["muted"],
                         line_spacing=1.22)
        if note:
            add_text(s, note, x, y + h - 0.6, w, 0.55, size=13, color=t["muted"], anchor="middle")
        return s

    def chart(self, title, kind, labels, values, insight="", unit="", insight_title="Nhận xét"):
        """Biểu đồ theo màu theme + thẻ nhận xét. kind: bar | barh | line | pie | donut."""
        t = self.t
        s, (x, y, w, h) = self.blank_slide(title)
        cw = 7.8
        add_box(s, x, y, cw, h, t["card"], line=t["border"], radius=0.05, shadow=True)
        img = make_chart(kind, labels, values, t, w_in=cw - 0.4, h_in=h - 0.4, unit=unit)
        s.shapes.add_picture(img, Inches(x + 0.2), Inches(y + 0.2), Inches(cw - 0.4), Inches(h - 0.4))
        ix, iw = x + cw + 0.35, w - cw - 0.35
        add_box(s, ix, y, iw, h, t["dark"], rounded=True, radius=0.05)
        add_icon(s, "idea", ix + 0.35, y + 0.4, 0.6, t["accent"])
        add_text(s, insight_title, ix + 0.3, y + 1.1, iw - 0.6, 0.6, size=20, color=t["on_dark"], font=TITLE_FONT)
        add_text(s, insight, ix + 0.3, y + 1.75, iw - 0.6, h - 2.05, size=17, min_size=12, color=t["on_dark_muted"],
                 line_spacing=1.28)
        return s

    def comparison(self, title, left, right):
        """So sánh 2 cột: left / right = (tiêu đề, list ý) hoặc (icon, tiêu đề, list ý)."""
        t = self.t
        s, (x, y, w, h) = self.blank_slide(title)
        gap = 0.4
        cw = (w - gap) / 2
        for i, side in enumerate((left, right)):
            icon, head, pts = (side if len(side) == 3 else (None,) + tuple(side))
            cx = x + i * (cw + gap)
            band = t["strong"] if i == 0 else t["accent"]
            add_box(s, cx, y, cw, h, t["card"], line=t["border"], radius=0.05, shadow=True)
            add_box(s, cx, y, cw, 0.95, band, radius=0.18)
            add_box(s, cx, y + 0.5, cw, 0.45, band, rounded=False)  # vuông góc dưới của dải màu
            if icon:
                add_icon(s, icon, cx + 0.3, y + 0.2, 0.55, "FFFFFF")
            add_text(s, head, cx + (0.95 if icon else 0.3), y + 0.12, cw - 1.2, 0.72, size=22, min_size=14,
                     color="FFFFFF", font=TITLE_FONT, anchor="middle")
            pts = pts if isinstance(pts, (list, tuple)) else [pts]
            n = max(len(pts), 1)
            rh = min(1.05, (h - 1.3) / n)
            for j, p in enumerate(pts[:6]):
                py = y + 1.2 + j * rh
                add_icon(s, "check-circle", cx + 0.32, py + 0.1, 0.42, band)
                add_text(s, str(p), cx + 0.9, py + 0.04, cw - 1.2, rh, size=21, min_size=12, color=t["text"])
        return s

    def image_grid(self, title, images, captions=None):
        """Lưới 2–4 ảnh bo góc kèm chú thích ngắn."""
        t = self.t
        s, (x, y, w, h) = self.blank_slide(title)
        imgs = list(images)[:4]
        caps = list(captions or [])
        n = max(len(imgs), 1)
        gap = 0.3
        cw = (w - gap * (n - 1)) / n
        ih = h - (0.7 if caps else 0)
        for i, img in enumerate(imgs):
            cx = x + i * (cw + gap)
            add_image(s, img, cx, y, cw, ih, t, radius=0.06)
            if i < len(caps):
                add_text(s, caps[i], cx, y + ih + 0.1, cw, 0.55, size=15, min_size=11, color=t["muted"], align="center")
        return s

    def timeline(self, title, steps):
        """Quy trình 3–5 bước: đường nối + vòng số + thẻ mô tả."""
        t = self.t
        s, (x, y, w, h) = self.blank_slide(title)
        items = [_item(p) for p in steps][:6]
        n = len(items)
        colw = w / n
        line_y = y + 0.55
        add_box(s, x + colw / 2, line_y - 0.02, w - colw, 0.04, t["border"], rounded=False)
        for i, (icon, head, body) in enumerate(items):
            cx = x + i * colw + colw / 2
            add_box(s, cx - 0.4, line_y - 0.4, 0.8, 0.8, t["accent"] if self.is_dark else t["strong"],
                    shape=MSO_SHAPE.OVAL, shadow=True)
            if icon:
                add_icon(s, icon, cx - 0.22, line_y - 0.22, 0.44, t["dark"] if self.is_dark else "FFFFFF")
            else:
                add_text(s, str(i + 1), cx - 0.4, line_y - 0.4, 0.8, 0.8, size=22, font=TITLE_FONT,
                         color=t["dark"] if self.is_dark else "FFFFFF", align="center", anchor="middle")
            self._card(s, x + i * colw + 0.1, line_y + 0.7, colw - 0.2, min(h - 1.25, 2.9), None, head, body, body_size=19)
        return s

    def statement(self, text, caption=""):
        """Một thông điệp lớn trên nền đậm."""
        t = self.t
        s = self._new(dark=True)
        self._decor(s, dark=True)
        add_box(s, 1.2, 1.7, 0.5, 0.07, t["accent"], radius=0.5)
        add_text(s, text, 1.15, 1.9, 10.6, 3.4, size=40, min_size=22, color=t["on_dark"], font=TITLE_FONT,
                 anchor="middle", line_spacing=1.1)
        if caption:
            add_text(s, caption, 1.15, 5.45, 10.4, 0.8, size=18, color=t["on_dark_muted"])
        self._footer(s, dark=True)
        return s

    def quote(self, text, author="", image=None):
        """Trích dẫn: ảnh tràn nền phủ tối, dấu ngoặc kép lớn."""
        t = self.t
        s = self._new(dark=True)
        if image:
            add_image(s, image, 0, 0, SLIDE_W, SLIDE_H, t)
            add_box(s, 0, 0, SLIDE_W, SLIDE_H, t["dark"], rounded=False, alpha=0.72)
        else:
            self._decor(s, dark=True)
        add_text(s, "“", 1.0, 0.6, 2, 2, size=150, color=t["accent"], font="Georgia")
        add_text(s, text, 1.5, 2.1, 10.3, 3.0, size=34, min_size=20, color=t["on_dark"], font=TITLE_FONT,
                 anchor="middle", line_spacing=1.15)
        if author:
            add_text(s, "— " + author, 1.5, 5.3, 10.3, 0.7, size=18, color=t["on_dark_muted"])
        return s

    def closing(self, title="Cảm ơn đã lắng nghe", subtitle="", image=None):
        t = self.t
        s = self._new(dark=True)
        if image:
            add_image(s, image, 0, 0, SLIDE_W, SLIDE_H, t)
            add_box(s, 0, 0, SLIDE_W, SLIDE_H, t["dark"], rounded=False, alpha=0.7)
        else:
            self._decor(s, dark=True)
            add_box(s, SLIDE_W / 2 - 3, SLIDE_H / 2 - 3, 6, 6, t["accent"], shape=MSO_SHAPE.OVAL, alpha=0.08)
        add_text(s, title, 1, 2.2, SLIDE_W - 2, 1.7, size=50, min_size=28, color=t["on_dark"], font=TITLE_FONT,
                 align="center", anchor="bottom")
        add_box(s, SLIDE_W / 2 - 0.3, 4.1, 0.6, 0.07, t["accent"], radius=0.5)
        if subtitle:
            add_text(s, subtitle, 1.5, 4.35, SLIDE_W - 3, 1.2, size=21, min_size=14, color=t["on_dark_muted"],
                     align="center")
        return s

    def save(self, filename="bai_thuyet_trinh.pptx") -> str:
        path = output_path(filename, ".pptx")
        self.prs.save(path)
        print(f"Đã lưu: {path} ({self.n} slide)")
        if WARNINGS:
            print("CẢNH BÁO thiết kế — hãy rút gọn các đoạn chữ sau rồi chạy lại để slide đẹp hơn:")
            for w in dict.fromkeys(WARNINGS):
                print("  -", w)
        return str(path)
